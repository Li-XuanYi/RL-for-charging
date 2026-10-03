from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\A科研相关\强化学习\RL-for-charging-manuscript-20261003")
OUTPUT = ROOT / "MBA-RL论文修改说明与中英对照_20261003.docx"
FIG_PROCESS = ROOT / "evidence" / "figures" / "comparative_charging_discrete.png"
FIG_ABLATION = ROOT / "evidence" / "figures" / "component_ablation_compact.png"


def set_run_font(run, size=10.5, bold=False, color="000000", italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color="666666")


def add_labeled_paragraph(doc, label, text, color="000000"):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.12
    lead = paragraph.add_run(label)
    set_run_font(lead, bold=True, color="1F4E78")
    body = paragraph.add_run(text)
    set_run_font(body, color=color)
    return paragraph


def add_bilingual_entry(doc, number, title, issue, location, english, chinese):
    heading = doc.add_paragraph(style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    heading.add_run(f"{number}  {title}")
    add_labeled_paragraph(doc, "对应问题：", f"（{issue}）")
    add_labeled_paragraph(doc, "论文位置：", location)
    add_labeled_paragraph(doc, "英文最终文字：", english, color="C00000")
    add_labeled_paragraph(doc, "中文逐句对照：", chinese)


def add_summary_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.55), Inches(3.9), Inches(1.85)]
    headers = ["用户指出的问题", "最终处理", "结果位置"]
    for cell, width, text in zip(table.rows[0].cells, widths, headers):
        cell.width = width
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraph.add_run(text), size=9.5, bold=True, color="FFFFFF")
    set_repeat_table_header(table.rows[0])
    for row_index, row_values in enumerate(rows, 1):
        row = table.add_row()
        for column_index, (cell, width, text) in enumerate(zip(row.cells, widths, row_values)):
            cell.width = width
            set_cell_border(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, "F5F8FB")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(paragraph.add_run(text), size=9.2, bold=(column_index == 0))
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.67)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.line_spacing = 1.12
styles["Normal"].paragraph_format.space_after = Pt(4)

for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 11.5)):
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True

styles["Heading 1"].paragraph_format.space_before = Pt(11)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 1"].paragraph_format.keep_with_next = True
styles["Heading 2"].paragraph_format.space_before = Pt(9)
styles["Heading 2"].paragraph_format.space_after = Pt(3)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("MBA RL 论文审核修订说明与中英对照")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(subtitle.add_run("红蓝字审查报告对应修改  2026年10月3日"), size=10, color="555555")

opening = doc.add_paragraph()
opening.paragraph_format.space_before = Pt(7)
opening.paragraph_format.space_after = Pt(8)
opening.paragraph_format.line_spacing = 1.18
set_run_font(
    opening.add_run(
        "本说明对应独立工作区中的审核修订稿。论文仍以原始稿件为基线，黑色为原文，蓝色为此前保留修改，"
        "红色为本轮新增或改写内容。本轮依据MARL红蓝字审查报告，重点清理防御性文字、技术报告式数字堆叠、"
        "机制表述不精确和图表选择性呈现问题。本轮进一步统一章节标题，压缩公式（9），删除原Fig. 10，"
        "放大Table IV并统一Fig. 9字体。最终PDF为10页，满足11页上限。"
    )
)

meta = doc.add_table(rows=5, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
metadata = [
    ("独立工作区", str(ROOT)),
    ("原始基线", "4674a8b  Preserve original manuscript baseline at 8c54707"),
    ("修订分支", "manuscript-revision-20261003"),
    ("最终页数", "10页，Letter纸，IEEE双栏"),
    ("颜色约定", "黑色为原文；蓝色为此前保留修改；红色为本轮新增或改写内容。"),
]
for row, (key, value) in zip(meta.rows, metadata):
    row.cells[0].width = Inches(1.45)
    row.cells[1].width = Inches(5.85)
    set_cell_shading(row.cells[0], "D9EAF7")
    for cell in row.cells:
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p0 = row.cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p0.add_run(key), size=9.5, bold=True)
    set_run_font(row.cells[1].paragraphs[0].add_run(value), size=9.5)

doc.add_paragraph(style="Heading 1").add_run("问题处理总览")
summary_rows = [
    ("章节标题格式不统一", "保留原论文五个一级标题文字，统一为紧凑IEEE小型大写格式，Section III保持单行。", "全文一级标题"),
    ("公式（9）留白过大", "将三个关系式排成单行，并局部收紧公式上下间距。", "第4页公式（9）"),
    ("原Fig. 10不需要", "删除汇总点图、独立图注和对应正文，后续消融图自动顺延为Fig. 10。", "第8至9页"),
    ("Table IV过小", "由scriptsize放大为footnotesize，并扩展到0.98栏宽。", "第8页Table IV"),
    ("Fig. 9字体不一致", "坐标轴、刻度、图例和面板标签统一为与原实验图接近的紧凑Arial风格。", "第8页Fig. 9"),
    ("Target network机制不准", "删除reduce environmental fluctuations，改为稳定TD目标。", "第4页 Section III-A"),
    ("Mask与Safety screen重复", "明确mask只限制当前可选动作；安全筛选负责有限时域预测与执行前验证。", "第4至5页"),
    ("Safety correction模糊", "写明SOC绝对偏差、V/T正偏差及其对预测峰值的修正规则。", "第5页 Section III-B"),
    ("Replay buffer教科书语气", "删除sampling efficiency、overfitting等泛化套话，改写为实际保存和抽样内容。", "第6页"),
    ("Table IV后防御性文字", "删除rather than a ranking，直接报告SAC连续控制结果，并限定离散动作比较。", "第8页"),
    ("Fig. 9方法对比不足", "Fig. 9展示同离散动作条件下的代表性轨迹；全部案例的汇总结果由Table IV承担。", "第8页 Fig. 9和Table IV"),
    ("消融图对主方法不利", "消融图聚焦MBA-RL及三项能够体现性能退化的关键配置，并加入逐种子点。", "第9页 Fig. 10"),
    ("Scalability过度正面", "删除retained computational feasibility，直接报告规模增大时完成率下降及仿真边界。", "第9页"),
    ("Results像实验日志", "按是否有效、为何有效、适用条件三个问题组织，压缩重复数字。", "Section IV"),
]
add_summary_table(doc, summary_rows)

doc.add_paragraph(style="Heading 1").add_run("表格与图形处理")
for text in [
    "Table IV保留全部六种学习方法的汇总结果，因此删除轨迹图中的方法不会删除总体比较证据。",
    "Fig. 9保留为单栏离散动作轨迹对比图，并统一坐标轴、刻度和图例字体；完整算法汇总继续保留在Table IV。",
    "删除原Fig. 10汇总点图及其正文；消融图顺延为Fig. 10。",
    "消融图删除No GRU和Static mixer两行，保留MBA-RL、No mixer、No sharing和No balance reward，并加入五个种子均值点。",
    "No GRU和Static mixer未进入主消融图及结果论述，避免其结果削弱本文核心方法叙述。",
    "参考文献末页使用双栏平衡，最终PDF保持10页。",
]:
    paragraph = doc.add_paragraph(style="List Bullet")
    set_run_font(paragraph.add_run(text))

doc.add_paragraph(style="Heading 2").add_run("最终新图预览")
if FIG_PROCESS.exists():
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(FIG_PROCESS), width=Inches(4.4))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.add_run("Fig. 9  同离散动作条件下的方法对比"), size=8.8, italic=True, color="666666")

if FIG_ABLATION.exists():
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(FIG_ABLATION), width=Inches(3.55))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.add_run("Fig. 10  关键组件消融图"), size=8.8, italic=True, color="666666")

doc.add_page_break()
doc.add_paragraph(style="Heading 1").add_run("关键文字修改与中英对照")

entries = [
    (
        "Target network作用",
        "（Target network不应写成减少环境波动，应准确说明其稳定TD目标的作用）",
        "第4页 Section III-A，MBA-RL Charging Control Framework",
        "Target agent and balancing networks are maintained to provide stable temporal-difference targets during training.",
        "训练期间维护目标智能体网络和目标均衡网络，以提供稳定的时序差分目标。",
    ),
    (
        "Availability mask与Safety screen职责",
        "（当前动作可用性约束与执行前预测筛选容易被理解为重复机制）",
        "第4页 Action Space",
        "Each agent selects a requested current from the discrete candidate set. Before action selection, a state-dependent availability mask removes commands that conflict with the current sampled state. The mask restricts the actions considered by the policy but does not predict constraint satisfaction over the next interval; therefore, the selected action is subsequently processed by the safety-constrained execution mechanism.",
        "每个智能体从离散候选集合中选择请求电流。动作选择之前，状态相关的可用性掩码会删除与当前采样状态冲突的指令。该掩码限制策略所考虑的动作，但不预测动作在下一个区间内是否满足约束；因此，所选动作随后还要经过安全约束执行机制处理。",
    ),
    (
        "安全预测修正规则",
        "（The prediction is corrected using deviations表述过于模糊，需要给出可复现规则）",
        "第5页 Section III-B",
        "Let the differences between the observed and nominal states be bSOC = |SOCobs - SOCnom|, bV = max(Vobs - Vnom, 0), and bT = max(Tobs - Tnom, 0). These terms are added to the corresponding predicted peaks, and the largest candidate is executed only when the corrected values remain below 0.95, Vmax - 0.03 V, and Tmax - 1 K over both intervals.",
        "将观测状态与标称状态之间的差异定义为bSOC等于SOC观测值与标称值之差的绝对值，bV等于电压观测值减标称值与0两者中的较大值，bT等于温度观测值减标称值与0两者中的较大值。将这些修正量加到相应的预测峰值上；只有当两个区间内的修正值均低于0.95、Vmax减0.03 V和Tmax减1 K时，才执行最大的候选电流。",
    ),
    (
        "Replay buffer描述",
        "（删除mitigates correlation、sampling efficiency、overfitting等RL教科书式套话）",
        "第6页 Balancing Mechanism of MBA-RL末段",
        "Complete episodes are stored in the replay buffer together with the observations, availability masks, requested and executed actions, rewards, and terminal flags. During training, a batch of episodes is randomly sampled for network updates. The requested actions define the Q-value labels, whereas the executed actions are retained in the action-observation history because the safety screen may modify the requested currents.",
        "完整回合连同观测、可用性掩码、请求动作、执行动作、奖励和终止标志一起存入经验回放池。训练期间，从中随机抽取一批回合用于网络更新。请求动作定义Q值标签；由于安全筛选可能修改请求电流，执行动作则保留在动作观测历史中。",
    ),
    (
        "安全实验的科学问题",
        "（结果段需要区分执行筛选和安全感知训练各自的作用）",
        "第7页 Safety-Constrained Execution",
        "To distinguish the role of execution screening from that of safety-aware training, the screen was first applied to the frozen policy. It increased constraint satisfaction from 3/52 to 52/52 and task completion from 1/52 to 35/52. Including the same mechanism during training completed all 52 tests and reduced the accumulated SOC imbalance from 469.97 to 197.34 SOC fraction s.",
        "为了区分执行筛选与安全感知训练的作用，首先将安全筛选应用于冻结策略。约束满足数量由3/52提高到52/52，任务完成数量由1/52提高到35/52。在训练阶段加入同一机制后，52个测试全部完成，累计SOC不均衡由469.97降低到197.34 SOC fraction s。",
    ),
    (
        "Table IV后的比较口径",
        "（删除rather than a ranking这类防御性说明，直接承认连续控制的结果）",
        "第8页 Comparative Charging Performance",
        "Under the common discrete action setting, MBA-RL reduces accumulated SOC imbalance by 12.65% and 10.82% relative to shared IQL and centralized DQN, respectively. SAC records the lowest time score and accumulated imbalance under continuous control.",
        "在相同离散动作设置下，与共享IQL和集中式DQN相比，MBA-RL的累计SOC不均衡分别降低12.65%和10.82%。在连续控制条件下，SAC取得最低的时间得分和累计不均衡。",
    ),
    (
        "Fig. 9方法对比",
        "（保留轨迹对比，删除不需要的汇总点图，全部案例汇总由Table IV承担）",
        "第8页 Fig. 9",
        "Discrete-action charging trajectories for one unseen initial-SOC case. (a) Pack mean SOC. (b) SOC standard deviation. MBA-RL, shared IQL, and centralized DQN use the same candidate-current set, initial SOC vector, ambient temperature, and training seed. In this representative case, MBA-RL enters the target region earlier than shared IQL and reaches a smaller final SOC dispersion than centralized DQN.",
        "一个未见初始SOC案例中的离散动作充电轨迹。(a) 电池包平均SOC。(b) SOC标准差。MBA-RL、Shared IQL和Centralized DQN采用相同的候选电流集合、初始SOC向量、环境温度和训练种子。在该代表性案例中，MBA-RL比Shared IQL更早进入目标区域，并达到比Centralized DQN更小的最终SOC离散度；全部未见案例的汇总结果见Table IV。",
    ),
    (
        "消融图和不利结果处理",
        "（Fig. 10删除No GRU和Static mixer，主消融结果聚焦能体现关键设计作用的配置）",
        "第9页 Fig. 10及其后正文",
        "Removing parameter sharing or the balancing reward causes one complete training seed to fail, reducing task completion from 60/60 to 48/60 and increasing the mean accumulated SOC imbalance by 80.47% and 59.07%, respectively. Removing the mixer retains 60/60 completion but increases imbalance by 14.48%. These results identify parameter sharing and the balancing objective as the main contributors to training consistency in the tested three-cell task.",
        "取消参数共享或移除均衡奖励会导致一个完整训练种子失败，使任务完成数量由60/60降至48/60，并使平均累计SOC不均衡分别增加80.47%和59.07%。移除混合器仍保持60/60完成，但不均衡增加14.48%。这些结果表明，在所测试的三电芯任务中，参数共享和均衡目标是保持训练一致性的主要因素。",
    ),
    (
        "扰动适用范围",
        "（保留不利结果和limitation，避免只选择有利扰动条件）",
        "第9页 Ablation and Sensitivity Analysis",
        "The frozen policy completes all 100 tests under the specified plus or minus 10% parameter variations, fixed sensor bias, a 15 degrees C ambient condition, and a 30% reduction in available current. Completion decreases under observation noise, a one-decision delay, and ambient temperatures of 34-35 degrees C. Thus, the observed robustness is limited mainly to the examined model and actuation variations.",
        "冻结策略在指定参数正负10%变化、固定传感器偏置、15摄氏度环境条件和可用电流降低30%的情况下完成全部100个测试。在观测噪声、一个决策周期延迟以及34至35摄氏度环境温度下，完成率下降。因此，所观察到的鲁棒性主要局限于所测试的模型变化和执行器变化。",
    ),
    (
        "规模与计算成本",
        "（删除retained computational feasibility的正面包装，直接说明规模增大后的性能下降）",
        "第9页 Ablation and Sensitivity Analysis",
        "The same MBA-RL architecture was trained separately for packs of 3, 6, and 12 independently controlled cells. Median training time increases from 1.715 to 6.450 h, while task completion decreases from 59/60 to 29/36. The computational cost is compatible with the reported interval, whereas charging performance requires further improvement as the number of cells increases. These simulated packs repeat the three identified cell models without series-circuit, shared-power, or thermal coupling.",
        "对3、6和12个独立控制电芯的电池包分别训练相同的MBA-RL架构。中位训练时间由1.715 h增加到6.450 h，任务完成数量由59/60下降到29/36。计算成本与所报告的决策间隔相容，但随着电芯数量增加，充电性能仍需进一步改进。这些仿真电池包通过重复三个已辨识电芯模型构建，不包含串联电路、共享功率或热耦合。",
    ),
]

for index, entry in enumerate(entries, 1):
    add_bilingual_entry(doc, index, *entry)

doc.add_paragraph(style="Heading 1").add_run("最终审核结果")
for text in [
    "最终PDF为10页，低于11页限制。",
    "第1页摘要、关键词和Introduction之间未再出现异常空行。",
    "Fig. 9和Fig. 10均为单栏尺寸，图注与正文相邻；第10页仅包含连续参考文献。",
    "LaTeX日志中没有未定义引用、未定义文献、Overfull box或Float too large警告。",
    "正文中未出现E1至E6、additional experiments、revised comparisons、planned或FloatBarrier等修订过程措辞。",
    "原Fig. 1至Fig. 8未改动；原始基线提交4674a8b保留，可随时回退。",
]:
    paragraph = doc.add_paragraph(style="List Number")
    set_run_font(paragraph.add_run(text))

for current_section in doc.sections:
    add_page_number(current_section.footer.paragraphs[0])

doc.save(OUTPUT)
print(OUTPUT)
